"""
excel_export.py
Tạo file Excel (.xlsx) chuyên nghiệp từ danh sách kết quả đọc MRZ hàng loạt.

Cột theo đúng yêu cầu: STT, Họ và tên, Ngày sinh, Giới tính, Số hộ chiếu, Quốc gia.
Bổ sung thêm 2 cột phục vụ nghiệp vụ thực tế (đối soát khi xử lý hàng loạt):
  - "Tên file ảnh": để đối chiếu ngược lại ảnh gốc, đặc biệt quan trọng với các dòng lỗi.
  - "Ghi chú": trạng thái OK / cảnh báo check digit / lỗi không đọc được ảnh.

LƯU Ý ĐỊNH DẠNG NGÀY SINH: cột "Ngày sinh" được ghi dưới dạng CHUỖI VĂN BẢN (text),
KHÔNG phải kiểu Date của Excel. Lý do: kiểu Date của Excel hiển thị khác nhau tùy theo
cài đặt vùng miền (Region) của từng máy tính (có máy hiện dd/mm/yyyy, có máy hiện
mm/dd/yyyy, gây nhầm lẫn ngày/tháng khi mở trên máy khác). Ghi dưới dạng text đảm bảo
ngày sinh LUÔN hiển thị đúng 1 định dạng cố định (dd/mm/yyyy) dù mở file ở bất kỳ máy
tính hay vùng miền nào.
"""

import io
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HEADER_FILL = PatternFill(start_color="1B2A4A", end_color="1B2A4A", fill_type="solid")
HEADER_FONT = Font(name="Arial", size=11, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=11)
ERROR_FONT = Font(name="Arial", size=11, color="B91C1C")
WARNING_FILL = PatternFill(start_color="FEF3C7", end_color="FEF3C7", fill_type="solid")
ERROR_FILL = PatternFill(start_color="FEE2E2", end_color="FEE2E2", fill_type="solid")
THIN = Side(style="thin", color="D9DCE1")
CELL_BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

SEX_LABELS = {"M": "Nam", "F": "Nữ", "X": "Không rõ"}

COLUMNS = [
    ("STT", 6),
    ("Tên file ảnh", 26),
    ("Họ và tên", 30),
    ("Ngày sinh", 14),
    ("Giới tính", 12),
    ("Số hộ chiếu", 16),
    ("Quốc gia", 20),
    ("Ghi chú", 34),
]


def _format_dob_as_text(iso_str: str) -> str:
    """Chuyển 'YYYY-MM-DD' -> 'DD/MM/YYYY' (chuỗi text thuần, không phải kiểu Date)."""
    try:
        d = datetime.strptime(iso_str, "%Y-%m-%d").date()
        return d.strftime("%d/%m/%Y")
    except (ValueError, TypeError):
        return iso_str or ""


def generate_excel(items: list) -> bytes:
    """
    items: list các dict, mỗi dict có dạng:
      { "filename": str, "success": bool, "data": {...MRZResult...} | None, "error": str | None }
    Trả về nội dung file .xlsx dạng bytes.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = "Ket qua doc MRZ"

    ws.append([c[0] for c in COLUMNS])
    for col_idx, (_, width) in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col_idx)].width = width
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = CELL_BORDER
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 22

    DOB_COL_INDEX = 4  # cột "Ngày sinh" - luôn ép định dạng text (xem ghi chú đầu file)

    for i, item in enumerate(items, start=1):
        filename = item.get("filename", "")
        row_idx = i + 1

        if item.get("success") and item.get("data"):
            d = item["data"]
            full_name = f"{d.get('surname', '')} {d.get('given_names', '')}".strip()
            birth_date_text = _format_dob_as_text(d.get("birth_date", ""))
            sex_label = SEX_LABELS.get(d.get("sex"), d.get("sex", ""))
            passport_no = d.get("passport_number", "")
            country = d.get("nationality_name") or d.get("issuing_country_name") or ""
            all_valid = d.get("all_checks_valid", False)
            note = "OK - toàn bộ check digit hợp lệ" if all_valid else \
                   "Cảnh báo: một số check digit KHÔNG khớp - nên kiểm tra lại ảnh gốc"
            row_values = [i, filename, full_name, birth_date_text, sex_label, passport_no, country, note]
            row_fill = None if all_valid else WARNING_FILL
            row_font = BODY_FONT
        else:
            error_msg = item.get("error", "Không đọc được MRZ từ ảnh")
            row_values = [i, filename, "-", "-", "-", "-", "-", f"LỖI: {error_msg}"]
            row_fill = ERROR_FILL
            row_font = ERROR_FONT

        ws.append(row_values)
        for col_idx in range(1, len(COLUMNS) + 1):
            cell = ws.cell(row=row_idx, column=col_idx)
            cell.font = row_font
            cell.border = CELL_BORDER
            if row_fill:
                cell.fill = row_fill
            if col_idx == 1:
                cell.alignment = Alignment(horizontal="center")
            if col_idx == DOB_COL_INDEX:
                # Ép kiểu ô là TEXT ('@') để Excel không tự diễn giải lại thành ngày/số
                cell.number_format = "@"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

